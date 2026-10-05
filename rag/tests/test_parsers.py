"""Tests for PDF and DOCX document parsers."""

import io
from pathlib import Path
import docx
from PIL import Image
import pymupdf
import pytest

from rag_service.parsers import parse_document
from rag_service.parsers.docx import extract_pages_from_docx
from rag_service.parsers.pdf import extract_pages_from_pdf


@pytest.fixture
def tiny_pdf_file(tmp_path: Path) -> Path:
    """Create a minimal valid 2-page PDF fixture."""
    pdf_path = tmp_path / "sample.pdf"
    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((50, 50), "First page: 5G RAN deployment specifications, massive MIMO beamforming configuration and cell capacity.")
    p2 = doc.new_page()
    p2.insert_text((50, 50), "Second page: Cloud core UPF network throughput benchmarks and user plane latency measurements.")
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def empty_pdf_file(tmp_path: Path) -> Path:
    """Create an empty 1-page PDF fixture with zero text."""
    pdf_path = tmp_path / "empty.pdf"
    doc = pymupdf.open()
    doc.new_page()  # blank page
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def scanned_image_pdf_file(tmp_path: Path) -> Path:
    """Create a 1-page PDF containing a raster image and zero digital text."""
    pdf_path = tmp_path / "scanned_image.pdf"
    img = Image.new("RGB", (100, 100), color=(240, 240, 240))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), stream=buf.getvalue())
    doc.save(str(pdf_path))
    doc.close()
    return pdf_path


@pytest.fixture
def tiny_docx_file(tmp_path: Path) -> Path:
    """Create a minimal valid DOCX fixture with headings and tables."""
    docx_path = tmp_path / "sample.docx"
    doc = docx.Document()
    doc.add_heading("Network Architecture", level=1)
    doc.add_paragraph("This document details the Standalone 5G core configuration.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parameter"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "MaxThroughput"
    table.cell(1, 1).text = "10Gbps"
    doc.save(str(docx_path))
    return docx_path


def test_pdf_parser_extracts_page_numbers_and_text(tiny_pdf_file: Path):
    """PDF parser must extract 1-indexed page tuples with clean text."""
    pages = extract_pages_from_pdf(tiny_pdf_file)
    assert len(pages) == 2
    assert pages[0][0] == 1
    assert "beamforming" in pages[0][1]
    assert pages[1][0] == 2
    assert "UPF" in pages[1][1]


def test_pdf_parser_rejects_empty_or_scanned_pdf(empty_pdf_file: Path, monkeypatch):
    """Empty PDF without OCR text must raise ValueError stating OCR is unavailable or found no text."""
    monkeypatch.setattr("rag_service.parsers.pdf._TESSERACT_AVAILABLE", False)
    with pytest.raises(ValueError) as excinfo:
        extract_pages_from_pdf(empty_pdf_file)
    assert "No extractable text found in PDF" in str(excinfo.value)
    assert "OCR is unavailable or found no text" in str(excinfo.value)


def test_pdf_parser_ocr_fallback_yields_text_when_available(scanned_image_pdf_file: Path, monkeypatch):
    """When OCR is available and page has no text layer, OCR extracts the page text."""
    monkeypatch.setattr("rag_service.parsers.pdf._TESSERACT_AVAILABLE", True)

    import pytesseract

    monkeypatch.setattr(
        pytesseract,
        "image_to_string",
        lambda img: "Extracted scanned content: Beamforming parameters for carrier aggregation.",
    )

    pages = extract_pages_from_pdf(scanned_image_pdf_file)
    assert len(pages) == 1
    assert pages[0][0] == 1
    assert "Beamforming parameters for carrier aggregation" in pages[0][1]


def test_pdf_parser_image_only_raises_error_when_ocr_unavailable(scanned_image_pdf_file: Path, monkeypatch):
    """When OCR is unavailable, an image-only PDF raises ValueError with the exact error message."""
    monkeypatch.setattr("rag_service.parsers.pdf._TESSERACT_AVAILABLE", False)
    with pytest.raises(ValueError) as excinfo:
        extract_pages_from_pdf(scanned_image_pdf_file)
    assert "No extractable text found in PDF" in str(excinfo.value)
    assert "OCR is unavailable or found no text" in str(excinfo.value)


def test_pdf_parser_normal_text_pdf_never_calls_ocr(tiny_pdf_file: Path, monkeypatch):
    """A PDF with rich digital text must extract text directly without calling OCR."""
    monkeypatch.setattr("rag_service.parsers.pdf._TESSERACT_AVAILABLE", True)

    ocr_called = False

    def mock_image_to_string(img):
        nonlocal ocr_called
        ocr_called = True
        return "Should not be called"

    import pytesseract

    monkeypatch.setattr(pytesseract, "image_to_string", mock_image_to_string)

    pages = extract_pages_from_pdf(tiny_pdf_file)
    assert len(pages) == 2
    assert "beamforming" in pages[0][1]
    assert "UPF" in pages[1][1]
    assert ocr_called is False


def test_docx_parser_extracts_headings_and_tables(tiny_docx_file: Path):
    """DOCX parser must extract paragraphs, headings, and table cells."""
    pages = extract_pages_from_docx(tiny_docx_file)
    assert len(pages) == 1
    page_num, text = pages[0]
    assert page_num == 1
    assert "Network Architecture" in text
    assert "Standalone 5G" in text
    assert "MaxThroughput" in text
    assert "10Gbps" in text


def test_parse_document_dispatcher_rejects_unsupported_extensions(tmp_path: Path):
    """parse_document must reject non-pdf/docx extensions with ValueError."""
    txt_file = tmp_path / "notes.txt"
    txt_file.write_text("Plain text notes.")
    with pytest.raises(ValueError) as excinfo:
        parse_document(txt_file)
    assert "Unsupported file format" in str(excinfo.value)
