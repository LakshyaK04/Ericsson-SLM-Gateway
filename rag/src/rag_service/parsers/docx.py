"""DOCX document parser using python-docx.

Extracts text, paragraphs, and heading styles from Microsoft Word documents.
"""

from pathlib import Path
from typing import List, Tuple, Union
import docx


def extract_pages_from_docx(file_path: Union[str, Path]) -> List[Tuple[int, str]]:
    """Extract structured text from a DOCX document.

    Returns a list of (page_number, text) tuples. For Word documents without
    fixed pagination, returns content grouped under page 1.

    Raises:
        ValueError: If no extractable text is present in the document.
    """
    doc = docx.Document(str(file_path))
    content_blocks: List[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            # Prefix headings to preserve structural cues
            style_name = para.style.name if para.style else ""
            if "Heading" in style_name or "Title" in style_name:
                content_blocks.append(f"\n## {text}\n")
            else:
                content_blocks.append(text)

    # Also extract any table contents
    for table in doc.tables:
        if not table.rows:
            continue
        first_row_cells = [" ".join(cell.text.split()).strip() for cell in table.rows[0].cells]
        has_headers = any(first_row_cells) and len(table.rows) > 1
        headers = [h if h else f"Col{i+1}" for i, h in enumerate(first_row_cells)]

        table_rows = []
        start_idx = 1 if has_headers else 0
        for row in table.rows[start_idx:]:
            row_parts = []
            for i, cell in enumerate(row.cells):
                cell_text = " ".join(cell.text.split()).strip()
                if not cell_text:
                    continue
                if has_headers:
                    h = headers[i] if i < len(headers) else f"Col{i+1}"
                    row_parts.append(f"{h}: {cell_text}")
                else:
                    row_parts.append(cell_text)
            if row_parts:
                table_rows.append(" | ".join(row_parts))
        if table_rows:
            content_blocks.append("\n" + "\n".join(table_rows) + "\n")

    full_text = "\n\n".join(content_blocks).strip()
    if not full_text:
        raise ValueError("No extractable text found in DOCX document; document appears empty.")

    return [(1, full_text)]
